import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Dict, Any

class DocxBroadsheetService:
    @staticmethod
    def generate_article_docx(article: Dict[str, Any]) -> io.BytesIO:
        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        is_arabic = article.get("language") == "ar"
        align_default = WD_ALIGN_PARAGRAPH.RIGHT if is_arabic else WD_ALIGN_PARAGRAPH.LEFT
        primary_font = "Dubai"

        # 1. Masthead Header
        masthead_p = doc.add_paragraph()
        masthead_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        platform_name = "صَحِّف" if is_arabic else "JOURNAL IT"
        run_name = masthead_p.add_run(f"{platform_name}\n")
        run_name.font.name = primary_font
        run_name.font.size = Pt(16)
        run_name.font.bold = True
        run_name.font.color.rgb = RGBColor(139, 0, 0) # Editorial Crimson

        run_sub = masthead_p.add_run("تقرير إخباري معتمد • نمط وكالات الأنباء" if is_arabic else "OFFICIAL NEWS WIRE REPORT")
        run_sub.font.name = primary_font
        run_sub.font.size = Pt(9)
        run_sub.font.color.rgb = RGBColor(100, 100, 100)

        # Rule line
        rule_p = doc.add_paragraph()
        rule_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rule_run = rule_p.add_run("══════════════════════════════════════════════════════════════════")
        rule_run.font.size = Pt(9)
        rule_run.font.color.rgb = RGBColor(180, 180, 180)
        rule_p.paragraph_format.space_after = Pt(14)

        # 2. Plain One-Shot Flowing Narrative Paragraphs (No titles, No subtitles)
        raw_text = article.get("report_text") or article.get("body_markdown") or ""
        paragraphs = raw_text.split("\n\n")
        for para in paragraphs:
            para = para.strip()
            # Remove any markdown header syntax if present
            if para.startswith("#"):
                para = para.lstrip("#").strip()
            if para.startswith("> "):
                para = para.replace("> ", "").strip()
            if not para:
                continue

            body_p = doc.add_paragraph()
            body_p.alignment = align_default
            b_run = body_p.add_run(para)
            b_run.font.name = primary_font
            b_run.font.size = Pt(12)
            b_run.font.color.rgb = RGBColor(25, 25, 25)
            body_p.paragraph_format.space_after = Pt(12)
            body_p.paragraph_format.line_spacing = 1.35

        footer_p = doc.add_paragraph()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        end_text = "— انتهى التقرير —" if is_arabic else "— END OF REPORT —"
        f_run = footer_p.add_run(end_text)
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(140, 140, 140)
        footer_p.paragraph_format.space_before = Pt(16)

        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream
